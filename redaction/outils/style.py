"""Mise en forme ECE du Word produit par pandoc (Times New Roman 12, double interligne, marges 2,54 cm, titres, tableaux)."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

import pathlib
B = pathlib.Path(__file__).resolve().parents[1] / "build"
d = Document(B / "brut.docx")
TNR = "Times New Roman"

def font(style, size=12, bold=None, italic=None):
    f = style.font; f.name = TNR; f.size = Pt(size); f.color.rgb = RGBColor(0, 0, 0)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), TNR)
    for t in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(t)) is not None: del rf.attrib[qn(t)]

# Marges 2,54 cm, A4
for s in d.sections:
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"): setattr(s, m, Cm(2.54))
    # numéro de page en bas au centre
    p = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag); r._r.append(e)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._r.append(e)
    r.font.name = TNR; r.font.size = Pt(11)

styles = d.styles
STY = {x.name: x for x in d.styles}
for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
    if name in [s.name for s in styles]:
        st = STY[name]; font(st, 12)
        pf = st.paragraph_format; pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        pf.space_before = Pt(0); pf.space_after = Pt(6)
        if name != "Compact": pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for lvl, size, align, indent in ((1, 14, WD_ALIGN_PARAGRAPH.CENTER, 0), (2, 12, WD_ALIGN_PARAGRAPH.LEFT, 0), (3, 12, WD_ALIGN_PARAGRAPH.LEFT, 1.27)):
    st = STY[f"Heading {lvl}"]; font(st, size, bold=True, italic=False)
    pf = st.paragraph_format; pf.alignment = align; pf.left_indent = Cm(indent)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE; pf.space_before = Pt(18 if lvl < 3 else 12); pf.space_after = Pt(12 if lvl < 3 else 6)
    pf.keep_with_next = True
if "Bibliography" in [s.name for s in styles]:
    st = STY["Bibliography"]; font(st, 12)
    pf = st.paragraph_format; pf.line_spacing_rule = WD_LINE_SPACING.SINGLE; pf.space_after = Pt(8)
    pf.left_indent = Cm(1.27); pf.first_line_indent = Cm(-1.27); pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
for name in ("Hyperlink",):
    if name in [s.name for s in styles]: STY[name].font.name = TNR

from docx.enum.style import WD_STYLE_TYPE
for nm in ("CaptionTable", "CaptionFigure"):
    if nm not in STY:
        st = d.styles.add_style(nm, WD_STYLE_TYPE.PARAGRAPH); st.base_style = STY["Normal"]; font(st, 11, bold=True); STY[nm] = st
for nm, size in (("TitlePage1", 18), ("TitlePage2", 13)):
    if nm in STY:
        st = STY[nm]; font(st, size, bold=(nm == "TitlePage1"))
        pf = st.paragraph_format; pf.alignment = WD_ALIGN_PARAGRAPH.CENTER; pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.space_before = Pt(40 if nm == "TitlePage1" else 14); pf.space_after = Pt(14)
# Word met à jour le sommaire et les listes à l'ouverture
_s = d.settings.element; _u = OxmlElement("w:updateFields"); _u.set(qn("w:val"), "true"); _s.append(_u)

# Titres de niveau 3 : terminés par un point (guide ECE)
for p in d.paragraphs:
    if p.style.name == "Heading 3" and p.runs:
        t = p.text.rstrip()
        if t and t[-1] not in ".?!:":
            p.runs[-1].text = p.runs[-1].text.rstrip() + "."
    # légendes (styles dédiés, pour les listes des tableaux et des figures) et sources : simple interligne
    m = re.match(r"^(Tableau|Figure) [A-Z0-9]+\.\d+ :", p.text)
    if m:
        p.style = STY["CaptionTable" if m.group(1) == "Tableau" else "CaptionFigure"]
    if m or p.text.startswith("Source :"):
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.keep_with_next = p.text.startswith("Tableau")
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if p.text.startswith("Figure") else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(6 if m else 2)
        p.paragraph_format.space_after = Pt(4 if p.text.startswith("Tableau") else 12)
        for r in p.runs: r.font.size = Pt(10 if p.text.startswith("Source") else 11)

# Tableaux : bordures, simple interligne, police 9-10
def borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "808080"); b.append(e)
    tblPr.append(b)
W_B1 = (2.7, 1.8, 3.6, 1.0, 1.6, 1.6, 1.4, 1.9)  # tableau B.1 (annexe) : colonnes assez larges pour ne pas couper les mots
for tbl in d.tables:
    borders(tbl)
    hdr = [c.text.strip() for c in tbl.rows[0].cells]
    if hdr[:3] == ["Extension", "Cible", "Comparaison"] and len(hdr) == len(W_B1):
        tbl.autofit = False
        grid = tbl._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), W_B1):
            gc.set(qn("w:w"), str(int(w * 567)))
        for row in tbl.rows:
            for cell, w in zip(row.cells, W_B1):
                cell.width = Cm(w)
    ncol = len(tbl.columns)
    size = 9 if ncol >= 5 else 10
    for ri, row in enumerate(tbl.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                p.paragraph_format.space_after = Pt(2); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.size = Pt(size); r.font.name = TNR
                    if ri == 0: r.font.bold = True
            if ri == 0:
                tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
                sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), "E7E6E6"); tcPr.append(sh)
# Sauts de page : un paragraphe vide contenant un saut de page crée une page blanche s'il tombe
# en haut de page. On le remplace par « saut de page avant » sur le paragraphe suivant.
for p in list(d.paragraphs):
    brs = p._p.xpath('.//w:br[@w:type="page"]')
    if brs and not p.text.strip():
        nxt = p._p.getnext()
        while nxt is not None and nxt.tag in (qn("w:bookmarkStart"), qn("w:bookmarkEnd")):
            nxt = nxt.getnext()
        if nxt is not None and nxt.tag == qn("w:p"):
            pPr = nxt.find(qn("w:pPr"))
            if pPr is None:
                pPr = OxmlElement("w:pPr"); nxt.insert(0, pPr)
            pb = OxmlElement("w:pageBreakBefore")
            pos = sum(1 for c in pPr if c.tag in (qn("w:pStyle"), qn("w:keepNext"), qn("w:keepLines")))
            pPr.insert(pos, pb)
            p._p.getparent().remove(p._p)
d.core_properties.author = "Abdellah Elyamine DALI BRAHAM"
d.core_properties.title = "Prédiction d'indicateurs des marchés financiers à partir des données de dépenses publiques ouvertes : une approche par machine learning"
d.save(B / "Memoire_DaliBraham.docx")
print("ok")

# paragraphes contenant une équation : simple interligne (affichage LibreOffice)
d = Document(B / "Memoire_DaliBraham.docx")
for p in d.paragraphs:
    if p._p.xpath(".//m:oMathPara") or p._p.xpath(".//m:oMath"):
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(6)
d.save(B / "Memoire_DaliBraham.docx")
print("équations ok")
